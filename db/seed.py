"""
OASI - Carga del Excel origen a PostgreSQL.

Lee "20260922 Levantamiento de Permisos Limpia.xlsx" (esperado en backend/data/)
y carga: empresas, proyectos, permisos, comites y permisos_comite. Requiere
que las migraciones ya se hayan aplicado (npm run db:migrate).

Uso:
    cd proyectos-backend
    .venv/bin/python db/seed.py [--file "data/otro_excel.xlsx"] [--dry-run]

CAMBIO DE FORMATO (22-09-2026): antes las columnas se leían por POSICIÓN
(números de columna hardcodeados) porque el Excel origen no tenía encabezados
de fila usables. La planilla nueva sí los tiene (fila 2 de "Proyectos" y
"Permisos"), así que ahora se leen por NOMBRE. Es más robusto: si alguien
reordena o agrega una columna en el Excel, el script no rompe en silencio
leyendo el dato equivocado, y si le sacan una columna que este script
necesita, avisa con el nombre en vez de leer una columna que no es.

Decisiones de limpieza (ver notas en cada función):
  - Fechas basura (año < 2000, 'S/I', '0', vacío, 'Por definir') -> NULL.
    fecha_registro_catastro además trae, en la planilla nueva, números de
    serie de Excel guardados como texto (ej. '46268') y fechas en texto
    'DD-MM-YYYY' (ej. '17-09-2026') — las dos formas se reconocen y se
    convierten; cualquier otra cosa rara queda en NULL, no se adivina.
  - Booleanos sucios ('Sí'/'Si'/'SI'/'si' -> true, 'No'/'no'/'NO' -> false,
    cualquier otra cosa (texto largo, número, vacío) -> el default que pida
    cada campo
  - Sectores: se unifican duplicados obvios ('Inmobiliarios'/'Inmobiliario',
    'Otros'/'Otro') y las dos variantes que el catálogo unificó
    ('Infraestructura' -> 'Infraestructura / Obras públicas',
    'Energía / Infraestructura' -> 'Energía'). Ver SECTOR_ALIAS.
  - "es_critico" viene 100% vacío en el Excel origen (las dos versiones) ->
    todos los permisos quedan con critico=false. Hay que revisarlos a mano en
    la app.
  - comite_registro trae texto tipo "Comité 10" o "Por definir" — se extrae el
    número con una expresión regular; "Por definir"/"por definir" (o
    cualquier otra cosa que no matchee) no vincula el permiso a ningún comité.
  - organismo -> ministerio: mapa fijo abajo. Ya NO se usa para insertar (ver
    abajo): queda como chequeo de que el Excel no traiga un organismo que el
    catálogo no conoce.

CATÁLOGOS (importante):
  regiones, sectores, etapas_proyecto, estados_permiso, tipologias,
  ministerios y organismos vienen PRE-CARGADOS por las migraciones, con ids
  explícitos y estables. Este script NO los inserta: busca cada valor del
  Excel por nombre y lo resuelve a su id. Si un valor del Excel no matchea
  ninguna fila del catálogo, el script ABORTA nombrando el valor, en vez de
  escribir un NULL silencioso que después nadie encuentra.
"""

import argparse
import datetime as dt
import os
import re
import sys

import openpyxl
import psycopg2

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_EXCEL = os.path.join(BACKEND_DIR, "data", "20260922 Levantamiento de Permisos Limpia.xlsx")

ORGANISMO_MINISTERIO = {
    "CONAF": "Ministerio de Agricultura",
    "SAG": "Ministerio de Agricultura",
    "BBNN": "Ministerio de Bienes Nacionales",
    "DGAC": "Ministerio de Defensa Nacional",
    "SSFFAA": "Ministerio de Defensa Nacional",
    "SEC": "Ministerio de Energía",
    "CEN": "Ministerio de Energía",
    "CMN": "Ministerio de las Culturas, las Artes y el Patrimonio",
    "MMA": "Ministerio de Medio Ambiente",
    "SEA": "Ministerio de Medio Ambiente",
    "SERNAGEOMIN": "Ministerio de Minería",
    "DGA": "Ministerio de Obras Públicas",
    "VIALIDAD": "Ministerio de Obras Públicas",
    "DOH": "Ministerio de Obras Públicas",
    "DIRECCION GENERAL DE CONCESIONES": "Ministerio de Obras Públicas",
    "SEREMI SALUD": "Ministerio de Salud",
    "MINVU": "Ministerio de Vivienda y Urbanismo",
    "DOM": "Municipalidades",
    "MTT": "Ministerio de Transportes y Telecomunicaciones",
}

# Variantes del Excel origen -> nombre canónico en el catálogo `sectores`.
# Las dos primeras son duplicados obvios (singular/plural); las dos últimas
# son las que unificó la migración a catálogos (0002_triggers_vistas_catalogos),
# y se repiten acá para que una carga desde cero dé exactamente el mismo
# resultado que una base ya migrada.
SECTOR_ALIAS = {
    "Inmobiliarios": "Inmobiliario",
    "Otros": "Otro",
    "Infraestructura": "Infraestructura / Obras públicas",
    "Energía / Infraestructura": "Energía",
}

# Estados de permiso tal como los escribe el Excel origen. Cualquier otra cosa
# (vacío, texto raro) cae en 'Pendiente', que es el default del schema.
ESTADOS_VALIDOS = ("Pendiente", "Resuelto", "Descartado", "Desistido")

# Nombres de columna esperados en la fila 2 (encabezado real) de cada hoja.
# Si el Excel no trae alguno de estos, el script avisa en vez de leer mal.
PROYECTOS_COLS = [
    "id_proyecto", "nombre_proyecto", "titular", "id_empresa", "empresa",
    "region", "sector", "tipologia", "inversion_mmusd", "empleo_construccion",
    "empleo_operacion", "estado_ambiental", "etapa_proyecto",
    "fecha_inicio_construccion", "fecha_inicio_operacion",
    "habilitantes_aprobado", "observaciones_oasi", "fecha_ingreso",
    "fecha_ultima_resolucion", "n_catastro", "incluido_en_catastro",
    "en_universo_permisos", "sigue_liberado_al_contactar",
    "listado_37_proyectos_liberados", "listado_97_proyectos_no_iniciados",
]

PERMISOS_COLS = [
    "id_permiso", "organismo", "nombre_permiso", "tipo_permiso",
    "nombre_permiso_estandar", "n_expediente", "es_critico", "que_habilita",
    "id_proyecto", "estado", "fecha_ingreso", "fecha_resolucion_estimada",
    "fecha_resolucion", "tipo_resolucion", "hito_tramitacion",
    "comite_registro", "es_permiso_habilitante_construccion",
    "incluido_catastro_hacienda", "observaciones", "n_catastro",
    "en_universo", "fecha_registro_catastro", "fecha_actualizacion",
    "quien_actualizo",
]

HEADER_ROW = 2
DATA_START_ROW = 3
PARAMETROS_COMITE_ROWS = range(8, 18)  # filas 8-17: Sesión (col B) / Fecha (col C)

COMITE_REGISTRO_RE = re.compile(r"comit[ée]\s*(\d+)", re.IGNORECASE)
EXCEL_EPOCH = dt.date(1899, 12, 30)  # base de las fechas-serie de Excel


# ----------------------------------------------------------------------------
# Limpieza
# ----------------------------------------------------------------------------

def clean_text(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s in ("S/I", "-", "N/A"):
        return None
    return s


def clean_date(v):
    """
    Acepta un datetime/date real (lo normal), un número de serie de Excel
    guardado como texto (dígitos puros, ej. '46268': pasa si cae entre 2015 y
    2035 aprox, para no confundirlo con otro número cualquiera), o una fecha
    en texto 'DD-MM-YYYY'. Cualquier otra cosa ('Por definir', '-', basura) es
    NULL: no se adivina.
    """
    if isinstance(v, dt.datetime):
        d = v.date()
    elif isinstance(v, dt.date):
        d = v
    elif isinstance(v, str):
        s = v.strip()
        if s.isdigit() and 42000 <= int(s) <= 50000:  # ~2015-01-01 a ~2036-11-21
            d = EXCEL_EPOCH + dt.timedelta(days=int(s))
        else:
            try:
                d = dt.datetime.strptime(s, "%d-%m-%Y").date()
            except ValueError:
                return None
    else:
        return None
    if d.year < 2000:  # fechas basura tipo 1900-03-29
        return None
    return d.isoformat()


def clean_bool(v, default=False):
    """
    Reconoce dos formatos, porque el Excel origen mezcla los dos según la
    columna: texto 'Sí'/'No' (la mayoría de los booleanos) y 1/0 numérico
    (ej. 'en_universo' en Permisos). Cualquier otra cosa (texto largo, vacío)
    cae en `default`.
    """
    if v is None:
        return default
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        if v == 1:
            return True
        if v == 0:
            return False
        return default
    s = str(v).strip().lower()
    if s in ("sí", "si", "s", "1"):
        return True
    if s in ("no", "n", "0"):
        return False
    return default  # texto largo, vacío, etc. -> default


def clean_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "" or s in ("S/I", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def clean_sector(v):
    s = clean_text(v)
    if s is None:
        return None
    return SECTOR_ALIAS.get(s, s)


def clean_int(v):
    n = clean_number(v)
    return int(n) if n is not None else None


def clean_comite_registro(v):
    """'Comité 10' -> 10. 'Por definir'/'por definir'/cualquier otra cosa -> None."""
    s = clean_text(v)
    if s is None:
        return None
    m = COMITE_REGISTRO_RE.search(s)
    return int(m.group(1)) if m else None


# ----------------------------------------------------------------------------
# Lectura del Excel
# ----------------------------------------------------------------------------

def find_last_real_row(ws, header_row, max_scan=2000, empty_streak_limit=50):
    last = header_row
    streak = 0
    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=max_scan, values_only=True), start=header_row + 1):
        if any(v is not None and str(v).strip() != "" for v in row):
            last = i
            streak = 0
        else:
            streak += 1
            if streak > empty_streak_limit and last > header_row:
                break
    return last


def read_sheet_by_header(ws, expected_cols, header_row=HEADER_ROW, data_start=DATA_START_ROW):
    """
    Lee una hoja usando el nombre de columna de `header_row`, no la posición.
    Devuelve una lista de dicts {nombre_columna: valor}. Aborta si falta
    alguna columna que el script necesita.
    """
    header = [c.value for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
    header_map = {}
    for idx, name in enumerate(header):
        if name is not None:
            header_map[str(name).strip()] = idx

    faltantes = [c for c in expected_cols if c not in header_map]
    if faltantes:
        raise CatalogoError(
            f"La hoja '{ws.title}' no tiene la(s) columna(s) esperada(s): {faltantes}\n"
            f"  Columnas encontradas en la fila {header_row}: {[h for h in header if h]}"
        )

    last_row = find_last_real_row(ws, header_row)
    out = []
    for row in ws.iter_rows(min_row=data_start, max_row=last_row, values_only=True):
        if all(v is None for v in row):
            continue
        out.append({name: row[idx] if idx < len(row) else None for name, idx in header_map.items()})
    return out


def read_comites(wb):
    ws = wb["Parámetros"]
    comites = {}
    for r in PARAMETROS_COMITE_ROWS:
        row = list(ws.iter_rows(min_row=r, max_row=r, values_only=True))[0]
        numero = row[1]  # col B
        fecha = row[2]   # col C
        if numero is not None and isinstance(fecha, (dt.datetime, dt.date)):
            comites[int(numero)] = clean_date(fecha)
    return comites


def read_proyectos(wb):
    ws = wb["Proyectos"]
    rows = read_sheet_by_header(ws, PROYECTOS_COLS)

    # Primera pasada: mapa nombre normalizado -> id_empresa real, de las filas
    # que sí traen los dos. Así, si CODELCO tiene id_empresa en una fila pero
    # no en otra, las dos quedan bajo el mismo id real en vez de la segunda
    # generando una fila de empresa nueva y separada para el mismo nombre.
    nombre_a_id = {}
    for row in rows:
        id_empresa_raw = clean_text(row.get("id_empresa"))
        nombre_empresa_raw = clean_text(row.get("empresa"))
        if id_empresa_raw and nombre_empresa_raw:
            nombre_a_id.setdefault(nombre_empresa_raw.strip().lower(), id_empresa_raw)

    proyectos = []
    empresas = {}  # id_excel -> nombre (primer nombre visto)
    for row in rows:
        id_excel = clean_text(row.get("id_proyecto"))
        if id_excel is None:
            continue
        id_empresa_raw = clean_text(row.get("id_empresa"))
        nombre_empresa_raw = clean_text(row.get("empresa"))
        if id_empresa_raw:
            empresa_id_excel = id_empresa_raw
            empresa_nombre = nombre_empresa_raw or id_empresa_raw
        elif nombre_empresa_raw:
            # Sin id_empresa pero con nombre (39/326 proyectos en la planilla
            # del 22-09): agrupar por nombre normalizado, no por el balde
            # compartido "SIN_EMPRESA" — si no, todos estos proyectos (de
            # empresas reales y distintas: CODELCO, Grenergy, Engie...)
            # terminaban bajo una sola fila de empresa, con el nombre de
            # cualquiera haya sido el primero que el script procesó y el
            # resto perdido. Si el nombre matchea una empresa que en OTRA
            # fila sí trae id_empresa, se usa ese id real; si no, la clave
            # sintética "NOMBRE:..." agrupa por nombre nomás — nunca se
            # guarda como id_excel (ver el insert más abajo).
            clave_nombre = nombre_empresa_raw.strip().lower()
            empresa_id_excel = nombre_a_id.get(clave_nombre, f"NOMBRE:{clave_nombre}")
            empresa_nombre = nombre_empresa_raw
        else:
            empresa_id_excel = "SIN_EMPRESA"
            empresa_nombre = "Sin empresa asignada"
        empresas.setdefault(empresa_id_excel, empresa_nombre)

        proyectos.append({
            "id_excel": id_excel,
            "nombre": clean_text(row.get("nombre_proyecto")) or id_excel,
            "titular": clean_text(row.get("titular")),
            "empresa_id_excel": empresa_id_excel,
            "region": clean_text(row.get("region")),
            "sector": clean_sector(row.get("sector")),
            "tipologia": clean_text(row.get("tipologia")),
            "inversion_mmusd": clean_number(row.get("inversion_mmusd")),
            "empleo_construccion": clean_int(row.get("empleo_construccion")),
            "empleo_operacion": clean_int(row.get("empleo_operacion")),
            "estado_ambiental": clean_text(row.get("estado_ambiental")),
            "etapa": clean_text(row.get("etapa_proyecto")),
            "fecha_inicio_construccion": clean_date(row.get("fecha_inicio_construccion")),
            "fecha_inicio_operacion": clean_date(row.get("fecha_inicio_operacion")),
            "habilitantes_aprobado": clean_bool(row.get("habilitantes_aprobado"), default=None),
            "fecha_ingreso": clean_date(row.get("fecha_ingreso")),
            "fecha_ultima_resolucion": clean_date(row.get("fecha_ultima_resolucion")),
            "observaciones_oasi": clean_text(row.get("observaciones_oasi")),
            "n_catastro": clean_int(row.get("n_catastro")),
            "incluido_en_catastro": clean_bool(row.get("incluido_en_catastro"), default=None),
            "en_universo_permisos": clean_bool(row.get("en_universo_permisos"), default=None),
            "sigue_liberado_al_contactar": clean_bool(row.get("sigue_liberado_al_contactar"), default=None),
            "listado_37_proyectos_liberados": clean_bool(row.get("listado_37_proyectos_liberados"), default=None),
            "listado_97_proyectos_no_iniciados": clean_bool(row.get("listado_97_proyectos_no_iniciados"), default=None),
        })
    return proyectos, empresas


def read_permisos(wb):
    ws = wb["Permisos"]
    rows = read_sheet_by_header(ws, PERMISOS_COLS)
    permisos = []
    for row in rows:
        proyecto_id_excel = clean_text(row.get("id_proyecto"))
        organismo = clean_text(row.get("organismo"))
        nombre_estandar = clean_text(row.get("nombre_permiso_estandar"))
        # Algunas filas no traen 'nombre_permiso' pero sí 'nombre_permiso_estandar' -> fallback.
        nombre = clean_text(row.get("nombre_permiso")) or nombre_estandar
        if proyecto_id_excel is None or organismo is None or nombre is None:
            continue  # sin proyecto/organismo/nombre no hay permiso valido

        estado = clean_text(row.get("estado")) or "Pendiente"
        if estado not in ESTADOS_VALIDOS:
            estado = "Pendiente"

        permisos.append({
            "id_excel": clean_text(row.get("id_permiso")),
            "organismo": organismo,
            "proyecto_id_excel": proyecto_id_excel,
            "nombre": nombre,
            "nombre_estandar": nombre_estandar,
            "tipo_permiso": clean_text(row.get("tipo_permiso")),
            "n_expediente": clean_text(row.get("n_expediente")),
            "critico": clean_bool(row.get("es_critico")),
            "que_habilita": clean_text(row.get("que_habilita")),
            "estado": estado,
            "fecha_ingreso": clean_date(row.get("fecha_ingreso")),
            "fecha_resolucion_estimada": clean_date(row.get("fecha_resolucion_estimada")),
            "fecha_resolucion": clean_date(row.get("fecha_resolucion")),
            "tipo_resolucion": clean_text(row.get("tipo_resolucion")),
            "hito_tramitacion": clean_text(row.get("hito_tramitacion")),
            "n_comite": clean_comite_registro(row.get("comite_registro")),
            "habilitante_construccion": clean_bool(row.get("es_permiso_habilitante_construccion")),
            "incluido_catastro_hacienda": clean_bool(row.get("incluido_catastro_hacienda"), default=None),
            "observaciones": clean_text(row.get("observaciones")),
            "n_catastro": clean_text(row.get("n_catastro")),
            "en_universo": clean_bool(row.get("en_universo"), default=None),
            "fecha_registro_catastro": clean_date(row.get("fecha_registro_catastro")),
            "fecha_actualizacion": clean_date(row.get("fecha_actualizacion")),
            "quien_actualizo": clean_text(row.get("quien_actualizo")),
        })
    return permisos


# ----------------------------------------------------------------------------
# Resolución de catálogos
#
# Los catálogos ya están en la base (los cargan las migraciones). Acá solo se los
# lee y se resuelve nombre -> id. Nada de INSERT: si el Excel trae un valor
# que el catálogo no tiene, es un dato a revisar, no una fila a crear.
# ----------------------------------------------------------------------------

class CatalogoError(Exception):
    """Un valor del Excel no matchea ninguna fila del catálogo, o falta una columna esperada."""


def cargar_catalogo(cur, tabla):
    """Devuelve {nombre: id} de una tabla de catálogo."""
    cur.execute(f"SELECT nombre, id FROM {tabla}")
    filas = cur.fetchall()
    if not filas:
        raise CatalogoError(
            f"El catálogo '{tabla}' está vacío. ¿Corriste npm run db:migrate? "
            f"Los catálogos se cargan con las migraciones, no con este script."
        )
    return {nombre: id_ for nombre, id_ in filas}


def resolver(catalogo, valor, etiqueta, faltantes):
    """
    nombre -> id. None pasa como None (columna opcional sin dato en el Excel).
    Un valor que no está en el catálogo se acumula en `faltantes` para poder
    reportarlos TODOS juntos y no morir en el primero.
    """
    if valor is None:
        return None
    if valor not in catalogo:
        faltantes.add((etiqueta, valor))
        return None
    return catalogo[valor]


def resolver_catalogos(cur, proyectos, permisos, organismos_en_permisos):
    """
    Resuelve, para cada fila leída del Excel, los ids de catálogo que
    necesitan los INSERT. Aborta con un mensaje que nombra cada valor no
    encontrado, en vez de dejar NULLs silenciosos.

    Muta `proyectos` y `permisos` agregándoles las claves *_id.
    """
    regiones = cargar_catalogo(cur, "regiones")
    sectores = cargar_catalogo(cur, "sectores")
    etapas = cargar_catalogo(cur, "etapas_proyecto")
    estados = cargar_catalogo(cur, "estados_permiso")
    organismos = cargar_catalogo(cur, "organismos")
    tipologias = cargar_catalogo(cur, "tipologias")

    faltantes = set()

    for p in proyectos:
        p["region_id"] = resolver(regiones, p["region"], "región", faltantes)
        p["sector_id"] = resolver(sectores, p["sector"], "sector", faltantes)
        p["etapa_id"] = resolver(etapas, p["etapa"], "etapa", faltantes)
        p["tipologia_id"] = resolver(tipologias, p["tipologia"], "tipología", faltantes)

    for p in permisos:
        p["estado_id"] = resolver(estados, p["estado"], "estado de permiso", faltantes)

    organismo_ids = {}
    for sigla in organismos_en_permisos:
        organismo_ids[sigla] = resolver(organismos, sigla, "organismo", faltantes)

    if faltantes:
        detalle = "\n".join(
            f"    - {etiqueta}: {valor!r}" for etiqueta, valor in sorted(faltantes)
        )
        raise CatalogoError(
            "Hay valores en el Excel que no existen en los catálogos de la base:\n"
            f"{detalle}\n"
            "  Corregí el Excel, o agregá el valor al catálogo en db/migrations/\n"
            "  (y a la migración correspondiente). No se escribió nada."
        )

    return organismo_ids


# ----------------------------------------------------------------------------
# Carga a Postgres
# ----------------------------------------------------------------------------

def load_db_config():
    env_path = os.path.join(BACKEND_DIR, ".env")
    cfg = {}
    if os.path.exists(env_path):
        with open(env_path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                k, v = line.split("=", 1)
                cfg[k.strip()] = v.strip()
    return cfg


def get_connection():
    cfg = load_db_config()
    return psycopg2.connect(
        host=os.environ.get("DB_HOST", cfg.get("DB_HOST", "localhost")),
        port=os.environ.get("DB_PORT", cfg.get("DB_PORT", "5432")),
        dbname=os.environ.get("DB_NAME", cfg.get("DB_NAME")),
        user=os.environ.get("DB_USER", cfg.get("DB_USER")),
        password=os.environ.get("DB_PASSWORD", cfg.get("DB_PASSWORD")),
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", default=DEFAULT_EXCEL)
    parser.add_argument("--dry-run", action="store_true", help="Solo parsear e imprimir conteos, no escribir en la BD")
    args = parser.parse_args()

    if not os.path.exists(args.file):
        print(f"No se encontró el Excel: {args.file}", file=sys.stderr)
        sys.exit(1)

    print(f"Leyendo {args.file} ...")
    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)

    try:
        comites = read_comites(wb)
        proyectos, empresas = read_proyectos(wb)
        permisos = read_permisos(wb)
    except CatalogoError as err:
        print(f"\nERROR: {err}", file=sys.stderr)
        sys.exit(1)

    organismos_en_permisos = sorted({p["organismo"] for p in permisos})
    organismos_sin_ministerio = [o for o in organismos_en_permisos if o not in ORGANISMO_MINISTERIO]

    print(f"  comites (calendario): {len(comites)}")
    print(f"  proyectos: {len(proyectos)}")
    print(f"  empresas unicas: {len(empresas)}")
    print(f"  permisos: {len(permisos)}")
    print(f"  organismos distintos en permisos: {len(organismos_en_permisos)}")
    if organismos_sin_ministerio:
        print(f"  ADVERTENCIA: organismos que el mapa fijo no conoce: {organismos_sin_ministerio}")

    proyectos_ids_excel = {p["id_excel"] for p in proyectos}
    permisos_sin_proyecto = [p for p in permisos if p["proyecto_id_excel"] not in proyectos_ids_excel]
    if permisos_sin_proyecto:
        print(f"  ADVERTENCIA: {len(permisos_sin_proyecto)} permisos referencian un proyecto que no existe, se omiten")

    permisos_sin_comite = sum(1 for p in permisos if p["n_comite"] is None)
    print(f"  permisos sin comité asignado ('Por definir' o vacío): {permisos_sin_comite}")

    if args.dry_run:
        # Sin tocar la base no se pueden resolver los catálogos, así que se
        # listan los valores distintos que trae el Excel para poder cotejarlos
        # a ojo contra los catálogos de la base antes de la carga real.
        print()
        print("Valores de vocabulario controlado que trae el Excel:")
        for etiqueta, clave, filas in (
            ("regiones", "region", proyectos),
            ("sectores (ya con alias aplicados)", "sector", proyectos),
            ("etapas", "etapa", proyectos),
            ("tipologías", "tipologia", proyectos),
            ("estados de permiso", "estado", permisos),
        ):
            valores = sorted({f[clave] for f in filas if f[clave] is not None})
            print(f"  {etiqueta} ({len(valores)}): {valores}")
        print()
        print("Dry-run: no se escribió nada en la base de datos.")
        return

    conn = get_connection()
    conn.autocommit = False
    cur = conn.cursor()

    try:
        # -- catálogos: se LEEN, no se insertan (vienen de las migraciones) --
        # Se resuelve todo antes del primer INSERT: si falta algún valor, el
        # script aborta acá y la transacción no llegó a escribir nada.
        organismo_ids = resolver_catalogos(cur, proyectos, permisos, organismos_en_permisos)

        # -- empresas --
        empresa_ids = {}
        for id_excel, nombre in empresas.items():
            # Ni "SIN_EMPRESA" ni una clave "NOMBRE:..." son un id real de la
            # planilla — las dos son sintéticas, agregadas acá para poder
            # agrupar filas sin id_empresa (ver read_proyectos).
            db_id_excel = None if id_excel == "SIN_EMPRESA" or id_excel.startswith("NOMBRE:") else id_excel
            cur.execute(
                "INSERT INTO empresas (id_excel, nombre) VALUES (%s, %s) RETURNING id",
                (db_id_excel, nombre),
            )
            empresa_ids[id_excel] = cur.fetchone()[0]

        # -- comites --
        comite_ids = {}
        for numero, fecha in sorted(comites.items()):
            cur.execute(
                "INSERT INTO comites (numero, fecha) VALUES (%s, %s) "
                "ON CONFLICT (numero) DO UPDATE SET fecha = EXCLUDED.fecha RETURNING id",
                (numero, fecha),
            )
            comite_ids[numero] = cur.fetchone()[0]

        # -- proyectos --
        proyecto_ids = {}
        for p in proyectos:
            cur.execute(
                """
                INSERT INTO proyectos (
                    id_excel, nombre, titular, empresa_id, region_id, sector_id,
                    tipologia_id, inversion_mmusd, empleo_construccion, empleo_operacion,
                    estado_ambiental, etapa_id, fecha_inicio_construccion,
                    fecha_inicio_operacion, habilitantes_aprobado, fecha_ingreso,
                    fecha_ultima_resolucion, observaciones_oasi, n_catastro,
                    incluido_en_catastro, en_universo_permisos, sigue_liberado_al_contactar,
                    listado_37_proyectos_liberados, listado_97_proyectos_no_iniciados
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    p["id_excel"], p["nombre"], p["titular"], empresa_ids[p["empresa_id_excel"]],
                    p["region_id"], p["sector_id"], p["tipologia_id"], p["inversion_mmusd"],
                    p["empleo_construccion"], p["empleo_operacion"], p["estado_ambiental"], p["etapa_id"],
                    p["fecha_inicio_construccion"], p["fecha_inicio_operacion"],
                    p["habilitantes_aprobado"], p["fecha_ingreso"], p["fecha_ultima_resolucion"],
                    p["observaciones_oasi"], p["n_catastro"], p["incluido_en_catastro"],
                    p["en_universo_permisos"], p["sigue_liberado_al_contactar"],
                    p["listado_37_proyectos_liberados"], p["listado_97_proyectos_no_iniciados"],
                ),
            )
            proyecto_ids[p["id_excel"]] = cur.fetchone()[0]

        # -- permisos + permisos_comite --
        permisos_cargados = 0
        permisos_comite_cargados = 0
        for p in permisos:
            proyecto_id = proyecto_ids.get(p["proyecto_id_excel"])
            organismo_id = organismo_ids.get(p["organismo"])
            if proyecto_id is None or organismo_id is None:
                continue

            cur.execute(
                """
                INSERT INTO permisos (
                    id_excel, proyecto_id, organismo_id, nombre, nombre_estandar, tipo_permiso,
                    n_expediente, critico, que_habilita, habilitante_construccion,
                    estado_id, fecha_ingreso, fecha_resolucion_estimada, fecha_resolucion,
                    tipo_resolucion, hito_tramitacion, incluido_catastro_hacienda,
                    n_catastro, observaciones, en_universo, fecha_registro_catastro,
                    fecha_actualizacion, quien_actualizo
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    p["id_excel"], proyecto_id, organismo_id, p["nombre"], p["nombre_estandar"], p["tipo_permiso"],
                    p["n_expediente"], p["critico"], p["que_habilita"], p["habilitante_construccion"],
                    p["estado_id"], p["fecha_ingreso"], p["fecha_resolucion_estimada"], p["fecha_resolucion"],
                    p["tipo_resolucion"], p["hito_tramitacion"], p["incluido_catastro_hacienda"],
                    p["n_catastro"], p["observaciones"], p["en_universo"], p["fecha_registro_catastro"],
                    p["fecha_actualizacion"], p["quien_actualizo"],
                ),
            )
            permiso_id = cur.fetchone()[0]
            permisos_cargados += 1

            if p["n_comite"] is not None and p["n_comite"] in comite_ids:
                cur.execute(
                    "INSERT INTO permisos_comite (permiso_id, comite_id) VALUES (%s, %s) "
                    "ON CONFLICT (permiso_id, comite_id) DO NOTHING",
                    (permiso_id, comite_ids[p["n_comite"]]),
                )
                permisos_comite_cargados += 1

        conn.commit()
        print()
        print("Carga completa:")
        print(f"  organismos referenciados (del catálogo): {len(organismo_ids)}")
        print(f"  empresas: {len(empresa_ids)}")
        print(f"  comites: {len(comite_ids)}")
        print(f"  proyectos: {len(proyecto_ids)}")
        print(f"  permisos: {permisos_cargados}")
        print(f"  permisos_comite: {permisos_comite_cargados}")

    except CatalogoError as err:
        # Error de datos, no un bug: se reporta legible y sin traceback.
        conn.rollback()
        print(f"\nERROR: {err}", file=sys.stderr)
        sys.exit(1)
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


if __name__ == "__main__":
    main()
